"""
Zet het ingevulde BiebMissie-invulsjabloon (.xlsx) om naar content/biebmissie.json.

Gebruik:
    python tools/xlsx_naar_biebmissie.py <invulsjabloon.xlsx> [content/biebmissie.json]

Vereist: openpyxl (pip install openpyxl)

Het script stopt met een duidelijke foutmelding als er iets ontbreekt of niet klopt
(bijv. een kluisantwoord met het verkeerde aantal letters).
"""
import json
import re
import sys
import unicodedata
from pathlib import Path

from openpyxl import load_workbook

TYPES = {"meerkeuze": "mcq", "open": "open", "koppelen": "match", "volgorde": "order"}
CLAIM_URL = "https://www.biebmissie.nl"
REQUIRED_STATIONS = 3


def normalize(value):
    text = unicodedata.normalize("NFD", str(value or ""))
    text = "".join(c for c in text if unicodedata.category(c) != "Mn")
    return re.sub(r"[^a-z0-9]", "", text.lower())


def cell(value):
    return str(value).strip() if value is not None else ""


def split_list(value):
    return [p.strip() for p in cell(value).split(";") if p.strip()]


def rows(ws):
    """Rijen als dicts op basis van de kopregel; lege rijen worden overgeslagen."""
    headers = [cell(c.value) for c in ws[1]]
    for idx, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(cell(v) for v in r):
            continue
        yield idx, dict(zip(headers, r))


def fail(errors):
    print("Er zijn problemen gevonden in het sjabloon:\n")
    for e in errors:
        print("  - " + e)
    sys.exit(1)


def convert(xlsx_path):
    wb = load_workbook(xlsx_path, data_only=False)
    errors = []

    # --- Boeken
    stations = []
    for idx, r in rows(wb["Boeken"]):
        sid = cell(r.get("Station-ID"))
        if not re.fullmatch(r"[A-Za-z0-9_-]{1,20}", sid):
            continue  # lege rij of toelichting onder de tabel
        station = {
            "id": sid,
            "title": cell(r.get("Titel")),
            "author": cell(r.get("Auteur")),
            "capacity": int(r.get("Exemplaren tegelijk") or 1),
            "copyLabel": cell(r.get("Hoe herkennen ze het juiste exemplaar?")),
            "teaser": cell(r.get("Teaser (1-2 zinnen, op het iPad-scherm)")),
            "questions": [],
        }
        if cell(r.get("BiebMissie-station?")).lower() == "ja":
            station["book"] = "youtube"
            station["priority"] = True
        if not station["title"]:
            errors.append(f"Boeken rij {idx}: titel ontbreekt")
        stations.append(station)
    by_id = {s["id"]: s for s in stations}

    # --- Vragen
    for idx, r in rows(wb["Vragen"]):
        sid = cell(r.get("Station-ID"))
        qtype_nl = cell(r.get("Type")).lower()
        prompt = cell(r.get("Vraag (zoals de kinderen hem lezen)"))
        if not qtype_nl and not prompt:
            continue  # lege voorbeeldrij
        where = f"Vragen rij {idx} ({sid})"
        if sid not in by_id:
            errors.append(f"{where}: onbekend Station-ID")
            continue
        if qtype_nl not in TYPES:
            errors.append(f"{where}: type moet meerkeuze, open, koppelen of volgorde zijn")
            continue
        if not prompt:
            errors.append(f"{where}: vraag ontbreekt")
            continue

        station = by_id[sid]
        q = {
            "id": f"{sid.lower()}-{len(station['questions']) + 1}",
            "type": TYPES[qtype_nl],
            "location": cell(r.get("Waar staat het? (hoofdstuk / bladzijde)")),
            "prompt": prompt,
            "hint": cell(r.get("Hint")),
        }
        correct = cell(r.get("Juist antwoord"))
        if not correct:
            errors.append(f"{where}: juist antwoord ontbreekt")
            continue

        if q["type"] == "mcq":
            wrong = [cell(r.get(f"Fout antwoord {n}")) for n in (1, 2, 3)]
            wrong = [w for w in wrong if w]
            if len(wrong) < 2:
                errors.append(f"{where}: meerkeuze heeft minstens 2 foute antwoorden nodig")
                continue
            # Juiste antwoord op een vaste maar per vraag verschillende plek
            pos = sum(map(ord, q["id"])) % (len(wrong) + 1)
            options = wrong[:pos] + [correct] + wrong[pos:]
            q["options"] = options
            q["correctIndex"] = pos
        elif q["type"] == "open":
            q["answers"] = [correct] + split_list(r.get("Ook goed (alleen bij open, gescheiden door ;)"))
        elif q["type"] == "match":
            pairs = []
            for part in split_list(correct):
                if "=" not in part:
                    errors.append(f"{where}: koppelen verwacht 'links = rechts; ...'")
                    break
                left, right = (p.strip() for p in part.split("=", 1))
                pairs.append({"left": left, "right": right})
            if len(pairs) != 4:
                errors.append(f"{where}: koppelen heeft precies 4 paren nodig (nu {len(pairs)})")
                continue
            q["pairs"] = pairs
        elif q["type"] == "order":
            items = split_list(correct)
            if len(items) != 4:
                errors.append(f"{where}: volgorde heeft precies 4 items nodig (nu {len(items)})")
                continue
            # Sla door elkaar op; correctOrder wijst terug naar de juiste volgorde
            shuffled_idx = [2, 0, 3, 1]
            q["items"] = [items[i] for i in shuffled_idx]
            q["correctOrder"] = [shuffled_idx.index(i) for i in range(4)]

        vault = cell(r.get("Kluisvraag"))
        if vault:
            q["vaultQuestion"] = int(float(vault))
        think = cell(r.get("Hersenspinsel (optioneel)"))
        if think:
            q["thinkQuestion"] = think
        station["questions"].append(q)

    for s in stations:
        if not s["questions"]:
            errors.append(f"Station {s['id']} heeft geen vragen")

    # --- Kluis (rijvolgorde = codeformulier)
    vault_rows = []
    answer_no = 0
    for n, (idx, r) in enumerate(rows(wb["Kluis"]), start=1):
        label = cell(r.get("Rij op het codeformulier"))
        if not label or label == "De code":
            continue
        answer = cell(r.get("Antwoord (invullen)"))
        length = int(r.get("Aantal vakjes") or 0)
        highlight = int(r.get("Gekleurd vakje nr") or 0)
        if not length:
            continue
        row = {"row": len(vault_rows) + 1, "label": label, "length": length, "highlight": highlight, "answer": answer}
        if label.lower() == "cadeautje":
            row["given"] = True
        else:
            answer_no += 1
            row["answerNo"] = answer_no
        if len(normalize(answer)) != length:
            errors.append(f"Kluis '{label}': {len(normalize(answer))} letters, het codeformulier heeft {length} vakjes")
        vault_rows.append(row)

    # Koppel kluisvragen (nummer 1-4) aan de juiste rij en controleer de antwoorden
    questions = [q for s in stations for q in s["questions"]]
    for row in vault_rows:
        if row.get("given"):
            continue
        linked = [q for q in questions if q.get("vaultQuestion") == row["answerNo"]]
        if len(linked) != 1:
            errors.append(f"Kluis '{row['label']}': verwacht precies 1 vraag met kluisvraag {row['answerNo']} (nu {len(linked)})")
            continue
        q = linked[0]
        q["vaultRow"] = row["row"]
        if q["type"] != "open":
            errors.append(f"Kluisvraag {row['answerNo']} moet type 'open' hebben")
        elif normalize(row["answer"]) not in [normalize(a) for a in q["answers"]]:
            errors.append(f"Kluisvraag {row['answerNo']}: antwoord op tabblad Vragen wijkt af van tabblad Kluis")
    for q in questions:
        q.pop("vaultQuestion", None)
    for row in vault_rows:
        row.pop("answerNo", None)

    if errors:
        fail(errors)

    return {
        "version": "1.0",
        "huntName": "BiebMissie MediaMasters",
        "requiredStations": REQUIRED_STATIONS,
        "stations": stations,
        "vault": {"claimUrl": CLAIM_URL, "rows": vault_rows},
    }


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(2)
    src = Path(sys.argv[1])
    dst = Path(sys.argv[2]) if len(sys.argv) > 2 else Path(__file__).resolve().parent.parent / "content" / "biebmissie.json"
    content = convert(src)
    dst.write_text(json.dumps(content, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    n = sum(len(s["questions"]) for s in content["stations"])
    print(f"OK: {len(content['stations'])} stations, {n} vragen -> {dst}")


if __name__ == "__main__":
    main()
